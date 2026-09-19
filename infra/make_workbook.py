"""
Build the sample contracts workbook.
================================================================================

    python -m infra.make_workbook

Writes `data/fiberops-contracts.xlsx`, the spreadsheet the code interpreter
reads. The file is committed, so you only need this script if you want to change
the data - it exists so the contents are reviewable in a diff, which a binary
xlsx is not.

The point of the workbook is that it holds what the *database does not*:
contractual SLAs, planned maintenance and the on-call rota. Answering "which
open alert is about to breach its SLA, and what does that cost?" needs both
sources, so the agent has to query Postgres over MCP *and* run pandas over this
file in the same turn. That is the whole reason both tools are attached.

Every `link_code` and site name below matches `infra/seed.sql` exactly. If you
change the seed data, change these too or the join silently returns nothing.
"""

from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = REPO_ROOT / "data" / "fiberops-contracts.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="6D3AED")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=11)

# --- SLA: the contract behind each link -------------------------------------
# max_resolution_hours is what makes this interesting: fiber_alerts.opened_at
# and resolved_at in Postgres can be measured against it.
SLA_HEADERS = [
    "link_code",
    "customer",
    "service_tier",
    "sla_availability_pct",
    "max_resolution_hours",
    "penalty_per_hour_brl",
    "contract_end",
]

SLA_ROWS = [
    ("LNK-SPO-RIO-01", "Banco Atlantico", "Platinum", 99.99, 4, 12500, date(2027, 12, 31)),
    ("LNK-SPO-BHZ-01", "Mineracao Serra Azul", "Gold", 99.95, 8, 6800, date(2027, 6, 30)),
    ("LNK-SPO-CWB-01", "Varejo Sul SA", "Gold", 99.95, 8, 5200, date(2026, 11, 30)),
    ("LNK-SPO-BSB-01", "Governo Federal - SERPRO", "Platinum", 99.99, 4, 18000, date(2028, 3, 31)),
    ("LNK-RIO-BHZ-01", "Petroquimica Guanabara", "Gold", 99.90, 12, 4300, date(2027, 1, 31)),
    ("LNK-RIO-BSB-01", "Midia Rio Broadcasting", "Silver", 99.50, 24, 1800, date(2026, 10, 31)),
    ("LNK-CWB-POA-01", "Agro Cooperativa Sul", "Gold", 99.90, 12, 5600, date(2027, 9, 30)),
    ("LNK-BHZ-BSB-01", "Universidade Federal MG", "Silver", 99.50, 24, 1200, date(2026, 12, 31)),
]

# --- Maintenance: planned work, which explains some of the alerts -----------
MAINTENANCE_HEADERS = [
    "window_id",
    "link_code",
    "starts_at",
    "ends_at",
    "work_type",
    "owner",
    "status",
]

MAINTENANCE_ROWS = [
    (
        "MW-2026-0141",
        "LNK-BHZ-BSB-01",
        datetime(2026, 9, 16, 1, 0),
        datetime(2026, 9, 16, 5, 0),
        "splice_repair",
        "Equipe Campo BH",
        "in_progress",
    ),
    (
        "MW-2026-0142",
        "LNK-SPO-BHZ-01",
        datetime(2026, 9, 19, 2, 0),
        datetime(2026, 9, 19, 6, 0),
        "amplifier_swap",
        "Equipe Campo SP",
        "scheduled",
    ),
    (
        "MW-2026-0143",
        "LNK-RIO-BSB-01",
        datetime(2026, 9, 21, 0, 0),
        datetime(2026, 9, 21, 4, 30),
        "otdr_survey",
        "Equipe Campo RJ",
        "scheduled",
    ),
    (
        "MW-2026-0138",
        "LNK-SPO-RIO-01",
        datetime(2026, 9, 9, 1, 0),
        datetime(2026, 9, 9, 3, 0),
        "connector_cleaning",
        "Equipe Campo SP",
        "completed",
    ),
    (
        "MW-2026-0139",
        "LNK-CWB-POA-01",
        datetime(2026, 9, 12, 2, 0),
        datetime(2026, 9, 12, 4, 0),
        "route_inspection",
        "Equipe Campo CWB",
        "completed",
    ),
]

# --- On-call: who to wake up, keyed by sites.name in Postgres ---------------
ONCALL_HEADERS = ["site", "shift", "engineer", "phone", "escalation_manager"]

ONCALL_ROWS = [
    ("Sao Paulo Central", "24x7", "Renata Alves", "+55 11 98812-4471", "Marcos Pinheiro"),
    ("Rio Downtown", "24x7", "Diego Fontes", "+55 21 99654-2210", "Marcos Pinheiro"),
    ("Belo Horizonte North", "business_hours", "Camila Duarte", "+55 31 98330-7712", "Ana Rezende"),
    ("Brasilia Core", "24x7", "Paulo Nogueira", "+55 61 99122-5508", "Ana Rezende"),
    ("Curitiba West", "business_hours", "Leticia Moraes", "+55 41 99881-3390", "Sergio Lima"),
    ("Porto Alegre South", "business_hours", "Bruno Tavares", "+55 51 99447-1265", "Sergio Lima"),
]


def _write_sheet(
    sheet: Worksheet,
    headers: list[str],
    rows: list[tuple],
    formats: dict[str, str] | None = None,
) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for row in rows:
        sheet.append(list(row))

    for column_index, header in enumerate(headers, start=1):
        letter = get_column_letter(column_index)
        width = max(len(header), *(len(str(row[column_index - 1])) for row in rows)) + 3
        sheet.column_dimensions[letter].width = min(width, 34)

        number_format = (formats or {}).get(header)
        for cell in sheet[letter][1:]:
            cell.font = BODY_FONT
            if number_format:
                cell.number_format = number_format

    sheet.freeze_panes = "A2"


def build() -> Path:
    workbook = Workbook()

    sla = workbook.active
    sla.title = "SLA"
    _write_sheet(
        sla,
        SLA_HEADERS,
        SLA_ROWS,
        {
            "sla_availability_pct": "0.00",
            "penalty_per_hour_brl": '"R$"#,##0',
            "contract_end": "yyyy-mm-dd",
        },
    )

    _write_sheet(
        workbook.create_sheet("Maintenance"),
        MAINTENANCE_HEADERS,
        MAINTENANCE_ROWS,
        {"starts_at": "yyyy-mm-dd hh:mm", "ends_at": "yyyy-mm-dd hh:mm"},
    )

    _write_sheet(workbook.create_sheet("OnCall"), ONCALL_HEADERS, ONCALL_ROWS)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> int:
    path = build()
    print(f"Wrote {path.relative_to(REPO_ROOT)}")
    print(f"  SLA          : {len(SLA_ROWS)} rows")
    print(f"  Maintenance  : {len(MAINTENANCE_ROWS)} rows")
    print(f"  OnCall       : {len(ONCALL_ROWS)} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
